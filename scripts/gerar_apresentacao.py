# -*- coding: utf-8 -*-
"""
Gera a apresentação PPTX (PT-BR):
"CRM Dental com Chat Inteligente usando Agentic RAG"
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ----------------------------------------------------------------------------
# Paleta / Design system
# ----------------------------------------------------------------------------
TEAL      = RGBColor(0x0F, 0x76, 0x6E)   # primária
TEAL_DK   = RGBColor(0x0B, 0x52, 0x4B)
CYAN      = RGBColor(0x06, 0xB6, 0xD4)   # acento
NAVY      = RGBColor(0x0F, 0x17, 0x2A)   # texto escuro / fundo capa
SLATE     = RGBColor(0x47, 0x55, 0x69)   # texto corpo
SLATE_LT  = RGBColor(0x94, 0xA3, 0xB8)
BG_LIGHT  = RGBColor(0xF8, 0xFA, 0xFC)
CARD      = RGBColor(0xFF, 0xFF, 0xFF)
CARD_ALT  = RGBColor(0xEC, 0xFE, 0xFF)   # cyan claro
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREEN     = RGBColor(0x16, 0xA3, 0x4A)
RED       = RGBColor(0xDC, 0x26, 0x26)
AMBER     = RGBColor(0xD9, 0x77, 0x06)
LINE      = RGBColor(0xCB, 0xD5, 0xE1)

FONT = "Segoe UI"

# 16:9
SW = Inches(13.333)
SH = Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def rect(s, x, y, w, h, color, line=None, line_w=Pt(1), shadow=False, rounded=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    fill(shp, color)
    if line is None:
        no_line(shp)
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    shp.shadow.inherit = False
    if shadow:
        _soft_shadow(shp)
    return shp


def _soft_shadow(shape):
    sp = shape._element.spPr
    existing = sp.find(qn('a:effectLst'))
    if existing is not None:
        sp.remove(existing)
    el = sp.makeelement(qn('a:effectLst'), {})
    sh = el.makeelement(qn('a:outerShdw'),
                        {'blurRad': '90000', 'dist': '40000',
                         'dir': '5400000', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '0F172A'})
    alpha = clr.makeelement(qn('a:alpha'), {'val': '18000'})
    clr.append(alpha)
    sh.append(clr)
    el.append(sh)
    sp.append(el)


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.05, wrap=True):
    """runs: list of paragraphs; each paragraph is list of (text, size, color, bold, italic)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, col, bold, ital) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(sz)
            r.font.color.rgb = col
            r.font.bold = bold
            r.font.italic = ital
            r.font.name = FONT
    return tb


def P(text, size, color, bold=False, italic=False):
    return [(text, size, color, bold, italic)]


def header(s, kicker, title, subtitle=None):
    """Standard content-slide header on light bg."""
    rect(s, 0, 0, Inches(0.22), SH, TEAL)            # left accent bar
    rect(s, Inches(0.6), Inches(0.55), Inches(0.55), Inches(0.12), CYAN)
    txt(s, Inches(0.6), Inches(0.72), Inches(11.5), Inches(0.4),
        [P(kicker.upper(), 13, TEAL, True)], space_after=0)
    txt(s, Inches(0.58), Inches(1.05), Inches(12.0), Inches(0.9),
        [P(title, 30, NAVY, True)], space_after=0)
    if subtitle:
        txt(s, Inches(0.6), Inches(1.72), Inches(12.0), Inches(0.5),
            [P(subtitle, 14, SLATE, False)], space_after=0)


def bullets(s, x, y, w, items, size=15, gap=10, color=SLATE,
            marker_color=TEAL, bold_lead=False):
    tb = s.shapes.add_textbox(x, y, w, Inches(4))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = "●  "
        r.font.size = Pt(size - 4)
        r.font.color.rgb = marker_color
        r.font.bold = True
        r.font.name = FONT
        if isinstance(it, tuple):
            lead, rest = it
            r2 = p.add_run()
            r2.text = lead
            r2.font.size = Pt(size)
            r2.font.bold = True
            r2.font.color.rgb = NAVY
            r2.font.name = FONT
            r3 = p.add_run()
            r3.text = rest
            r3.font.size = Pt(size)
            r3.font.color.rgb = color
            r3.font.name = FONT
        else:
            r2 = p.add_run()
            r2.text = it
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.bold = bold_lead
            r2.font.name = FONT
    return tb


def card(s, x, y, w, h, title, desc, accent=TEAL, icon=None):
    c = rect(s, x, y, w, h, CARD, shadow=True, rounded=True)
    rect(s, x, y, Inches(0.10), h, accent, rounded=False)
    yy = y + Inches(0.18)
    if icon:
        badge = rect(s, x + Inches(0.28), yy, Inches(0.5), Inches(0.5), accent, rounded=True)
        tfb = badge.text_frame
        tfb.word_wrap = False
        pp = tfb.paragraphs[0]
        pp.alignment = PP_ALIGN.CENTER
        rr = pp.add_run(); rr.text = icon; rr.font.size = Pt(18); rr.font.bold = True
        rr.font.color.rgb = WHITE; rr.font.name = FONT
        badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    txt(s, x + Inches(0.28), yy + (Inches(0.6) if icon else Inches(0)),
        w - Inches(0.5), Inches(0.4),
        [P(title, 15, NAVY, True)], space_after=0)
    txt(s, x + Inches(0.28), yy + (Inches(1.0) if icon else Inches(0.42)),
        w - Inches(0.5), h - Inches(1.2),
        [P(desc, 12.5, SLATE)], space_after=0, line_spacing=1.05)
    return c


def connector(s, x1, y1, x2, y2, color=SLATE_LT, width=Pt(2.0), arrow=True):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = width
    if arrow:
        ln = cn.line._get_or_add_ln()
        tail = ln.makeelement(qn('a:tailEnd'),
                              {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(tail)
    cn.shadow.inherit = False
    return cn


def pill(s, x, y, w, h, text, fillc, textc, size=12.5, bold=True):
    p = rect(s, x, y, w, h, fillc, rounded=True, shadow=True)
    tf = p.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    par = tf.paragraphs[0]
    par.alignment = PP_ALIGN.CENTER
    r = par.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = textc; r.font.name = FONT
    return p


def table(s, x, y, w, headers, rows, col_widths=None, header_fill=TEAL,
          row_h=Inches(0.5), head_h=Inches(0.5), fs=12.5, head_fs=13,
          zebra=True, cell_colors=None, align_first_left=True):
    nrows = len(rows) + 1
    ncols = len(headers)
    total_h = head_h + row_h * len(rows)
    gtbl = s.shapes.add_table(nrows, ncols, x, y, w, total_h)
    tbl = gtbl.table
    # disable default style banding
    tbl.first_row = False
    tbl.horz_banding = False
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = cw
    tbl.rows[0].height = head_h
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        tfc = cell.text_frame; tfc.word_wrap = True
        p = tfc.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if (align_first_left and j == 0) else PP_ALIGN.CENTER
        r = p.add_run(); r.text = h; r.font.bold = True; r.font.size = Pt(head_fs)
        r.font.color.rgb = WHITE; r.font.name = FONT
    for i, row in enumerate(rows):
        tbl.rows[i + 1].height = row_h
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            base = CARD if (i % 2 == 0) else BG_LIGHT
            if cell_colors and cell_colors.get((i, j)):
                base = cell_colors[(i, j)]
            cell.fill.solid(); cell.fill.fore_color.rgb = base
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            tfc = cell.text_frame; tfc.word_wrap = True
            p = tfc.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (align_first_left and j == 0) else PP_ALIGN.CENTER
            r = p.add_run(); r.text = val
            r.font.size = Pt(fs)
            r.font.bold = (j == 0)
            r.font.color.rgb = NAVY if (j == 0) else SLATE
            r.font.name = FONT
    return gtbl


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def footer(s, idx):
    txt(s, Inches(11.2), Inches(7.02), Inches(2.0), Inches(0.35),
        [P(f"CRM Dental · Agentic RAG     {idx:02d}", 9, SLATE_LT, False)],
        align=PP_ALIGN.RIGHT, space_after=0)


# ============================================================================
# SLIDE 1 — Capa
# ============================================================================
s = slide()
bg(s, NAVY)
# decorative shapes
d = rect(s, Inches(8.6), Inches(-1.5), Inches(7), Inches(7), TEAL_DK, rounded=True)
d.rotation = 25
d2 = rect(s, Inches(10.2), Inches(3.6), Inches(6), Inches(6), TEAL, rounded=True)
d2.rotation = 25
d3 = rect(s, Inches(11.0), Inches(0.2), Inches(2.2), Inches(2.2), CYAN, rounded=True)
d3.rotation = 25
# accent tag
rect(s, Inches(0.9), Inches(1.7), Inches(0.6), Inches(0.13), CYAN)
txt(s, Inches(0.9), Inches(1.95), Inches(7), Inches(0.5),
    [P("PROJETO ACADÊMICO · MVP", 14, CYAN, True)], space_after=0)
txt(s, Inches(0.86), Inches(2.45), Inches(8.3), Inches(2.4),
    [P("CRM Dental com Chat", 46, WHITE, True),
     P("Inteligente usando ", 46, WHITE, True),
     P("Agentic RAG", 46, CYAN, True)],
    space_after=2, line_spacing=1.02)
rect(s, Inches(0.92), Inches(5.15), Inches(3.4), Inches(0.04), SLATE)
txt(s, Inches(0.9), Inches(5.35), Inches(8), Inches(0.8),
    [P("Arquitetura técnica atual: CRM, RAG local, orquestração agentic e observabilidade", 18, SLATE_LT, False)],
    space_after=0)
txt(s, Inches(0.9), Inches(6.55), Inches(9), Inches(0.5),
    [P("Next.js · NestJS · Postgres · ChromaDB · Ollama · Agent traces",
       13, SLATE_LT, False)], space_after=0)
notes(s, "Boa tarde. Vou apresentar um MVP acadêmico de um CRM para clínicas "
         "odontológicas que incorpora um chat inteligente. O diferencial é a "
         "implementação atual de um RAG por paciente evoluído para um fluxo agentic, "
         "com orquestração no backend, recuperação vetorial, ferramentas de "
         "agendamento, validação de evidência e um registro persistente de cada "
         "execução para estudo e auditoria.")

# ============================================================================
# SLIDE 2 — Problema
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Contexto", "O problema nas clínicas odontológicas",
       "Dados clínicos existem, mas são pouco aproveitados no dia a dia.")
cards = [
    ("Histórico disperso", "Informações do paciente espalhadas em fichas, PDFs, anotações e sistemas diferentes.", "🗂"),
    ("Busca lenta", "Encontrar a última queixa ou um procedimento exige folhear vários registros.", "🔎"),
    ("Processos manuais", "Agendamentos e acompanhamentos dependem de trabalho repetitivo da equipe.", "🗓"),
    ("Dados subutilizados", "O conhecimento clínico acumulado raramente vira apoio à decisão.", "💤"),
]
cw = Inches(2.92); gap = Inches(0.18); x0 = Inches(0.6); y0 = Inches(2.45); ch = Inches(3.4)
for i, (t, d, ic) in enumerate(cards):
    card(s, x0 + i * (cw + gap), y0, cw, ch, t, d, accent=[TEAL, CYAN, TEAL, CYAN][i], icon=ic)
footer(s, 2)
notes(s, "O ponto de partida é prático. Em clínicas odontológicas, o histórico do "
         "paciente está fragmentado entre fichas de papel, PDFs e anotações livres. "
         "Isso torna a busca por informações lenta — por exemplo, lembrar qual foi a "
         "última queixa de dor. Além disso, agendamentos e acompanhamentos são manuais "
         "e os dados clínicos acumulados quase nunca são usados de forma inteligente. "
         "O objetivo do projeto é atacar exatamente essas quatro dores.")

# ============================================================================
# SLIDE 3 — Proposta do MVP
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Solução", "Proposta do MVP: um CRM Dental com chat",
       "Centraliza os dados do paciente e os torna consultáveis em linguagem natural.")
feats = [
    ("Cadastro de pacientes", "Dados demográficos e de contato centralizados.", "👤"),
    ("Anamnese estruturada", "Alergias, medicações, histórico e consentimento em Postgres.", "📋"),
    ("Agendamentos", "Status machine, auto-agendamento e aprovação pela clínica.", "📅"),
    ("Documentos + ingestão", "PDF/TXT/JSON/HTML viram chunks indexados no ChromaDB.", "📎"),
    ("Chat agentic", "Consulta evidências, usa ferramentas e mostra o trace da execução.", "💬"),
]
cw = Inches(3.85); ch = Inches(1.65); gx = Inches(0.22); gy = Inches(0.22)
positions = [(0,0),(1,0),(2,0),(0,1),(1,1)]
x0 = Inches(0.6); y0 = Inches(2.45)
for (t, d, ic), (cx, cy) in zip(feats, positions):
    card(s, x0 + cx*(cw+gx), y0 + cy*(ch+gy), cw, ch, t, d,
         accent=TEAL if cx % 2 == 0 else CYAN, icon=ic)
# highlight chat
hl = rect(s, x0 + 2*(cw+gx), y0 + 1*(ch+gy), cw, ch, TEAL, rounded=True, shadow=True)
tf = hl.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.25)
p = tf.paragraphs[0]; r = p.add_run(); r.text = "🌟  Diferencial do projeto"
r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "O chat com Agentic RAG conecta tudo isso."
r2.font.size = Pt(12.5); r2.font.color.rgb = CARD_ALT; r2.font.name = FONT
footer(s, 3)
notes(s, "A proposta é um CRM Dental enxuto, mas completo no essencial: cadastro de "
         "pacientes, anamnese estruturada, agendamentos, documentos e ingestão RAG. "
         "O diferencial atual é o chat agentic: além de consultar informações em "
         "linguagem natural, ele pode usar ferramentas de agendamento e expõe o curso "
         "de ações em um registro persistido para estudo.")

# ============================================================================
# SLIDE 4 — RAG tradicional (fluxo)
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Fundamento", "RAG tradicional: busca + geração",
       "Retrieval-Augmented Generation: o LLM responde com base em contexto recuperado.")
steps = [("Usuário", "👤", CYAN), ("Pergunta", "❓", CYAN),
         ("Busca vetorial", "🧭", TEAL), ("Contexto", "📄", TEAL),
         ("LLM", "🧠", TEAL_DK), ("Resposta", "✅", GREEN)]
n = len(steps)
bw = Inches(1.75); bh = Inches(1.25); y = Inches(3.3)
total = n * bw + (n - 1) * Inches(0.30)
x0 = (SW - total) / 2
for i, (t, ic, col) in enumerate(steps):
    x = x0 + i * (bw + Inches(0.30))
    b = rect(s, x, y, bw, bh, CARD, rounded=True, shadow=True)
    rect(s, x, y, bw, Inches(0.12), col)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = ic; r.font.size = Pt(24); r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = t; r2.font.size = Pt(13); r2.font.bold = True
    r2.font.color.rgb = NAVY; r2.font.name = FONT
    if i < n - 1:
        connector(s, x + bw, y + bh/2, x + bw + Inches(0.30), y + bh/2, color=SLATE_LT)
txt(s, Inches(0.6), Inches(5.4), Inches(12), Inches(0.8),
    [P("A pergunta vira um vetor (embedding), o sistema busca os trechos mais "
       "semelhantes no banco vetorial e injeta esse contexto no prompt do LLM, "
       "que então gera a resposta.", 14, SLATE)], space_after=0)
footer(s, 4)
notes(s, "Antes de falar de agentes, é importante entender o RAG tradicional. "
         "O fluxo é linear: a pergunta do usuário é convertida em um vetor numérico "
         "(embedding); fazemos uma busca por similaridade num banco vetorial para "
         "recuperar os trechos mais relevantes; esse contexto é inserido no prompt; "
         "e o LLM gera a resposta. É simples e eficaz, mas tem limitações importantes — "
         "que vejo no próximo slide.")

# ============================================================================
# SLIDE 5 — Limitações do RAG tradicional
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Motivação", "Limitações do RAG tradicional",
       "Um único passo de busca não dá conta de perguntas complexas.")
lims = [
    ("Recupera uma única vez", "Sem nova tentativa se o contexto vier fraco.", "🔁"),
    ("Contexto pouco relevante", "Pode trazer trechos que não respondem à pergunta.", "🎯"),
    ("Não planeja tarefas", "Não decompõe perguntas em etapas.", "🧩"),
    ("Não escolhe ferramentas", "Não consulta banco estruturado nem APIs.", "🛠"),
    ("Sofre com perguntas complexas", "Perguntas multi-etapas ficam incompletas.", "🌀"),
    ("Não valida a resposta", "Sem checagem de fundamentação nas fontes.", "⚠"),
]
cw = Inches(3.85); ch = Inches(1.55); gx = Inches(0.22); gy = Inches(0.20)
x0 = Inches(0.6); y0 = Inches(2.35)
for i, (t, d, ic) in enumerate(lims):
    cx = i % 3; cy = i // 3
    card(s, x0 + cx*(cw+gx), y0 + cy*(ch+gy), cw, ch, t, d, accent=AMBER, icon=ic)
footer(s, 5)
notes(s, "O RAG tradicional recupera contexto uma única vez: se a busca falhar, não "
         "há segunda tentativa. Ele pode trazer documentos pouco relevantes, não "
         "planeja a tarefa, não escolhe ferramentas (como consultar o banco relacional "
         "ou uma API de calendário) e tem dificuldade com perguntas que exigem várias "
         "etapas. Por fim, não valida se a resposta está realmente fundamentada nas "
         "fontes. Essas limitações motivam a evolução para Agentic RAG.")

# ============================================================================
# SLIDE 6 — Evolução para Agentic RAG
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Evolução", "De RAG para Agentic RAG",
       "Agentes especializados que planejam, decidem e cooperam.")
agents = [
    ("Orquestrador", "Controla a máquina de estados", "🧭", TEAL),
    ("Roteador", "Classifica intenção e custo esperado", "🧠", CYAN),
    ("Planner/Rewriter", "Decompõe e normaliza consultas", "🧩", TEAL),
    ("Retriever", "Combina ChromaDB e fallback Postgres", "🔎", CYAN),
    ("Evaluator/Verifier", "Avalia contexto e fundamentação", "✔", TEAL),
    ("Tool Registry", "Executa ações com preview → commit", "🛠", CYAN),
]
cw = Inches(3.85); ch = Inches(1.5); gx = Inches(0.22); gy = Inches(0.20)
x0 = Inches(0.6); y0 = Inches(2.35)
for i, (t, d, ic, col) in enumerate(agents):
    cx = i % 3; cy = i // 3
    card(s, x0 + cx*(cw+gx), y0 + cy*(ch+gy), cw, ch, t, d, accent=col, icon=ic)
footer(s, 6)
notes(s, "Na implementação atual, Agentic RAG não significa agentes autônomos soltos. "
         "Significa um conjunto de serviços especializados coordenados por uma "
         "máquina de estados determinística no NestJS. O roteador classifica a "
         "intenção, o planner e o rewriter preparam consultas, o retriever combina "
         "ChromaDB com fallback Postgres, o evaluator e o verifier controlam a "
         "qualidade da evidência, e o Tool Registry executa ações somente com "
         "confirmação humana.")

# ============================================================================
# SLIDE 7 — Arquitetura geral
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Arquitetura", "Visão geral da arquitetura",
       "Frontend → API NestJS → Orquestrador → RAG primitives / Postgres → Resposta.")

def archbox(x, y, w, h, title, sub, color, tcolor=WHITE):
    b = rect(s, x, y, w, h, color, rounded=True, shadow=True)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.size = Pt(13.5); r.font.bold = True
    r.font.color.rgb = tcolor; r.font.name = FONT
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(9.5)
        r2.font.color.rgb = tcolor; r2.font.name = FONT
    return b

yc = Inches(2.55)
fbox = archbox(Inches(0.6), yc, Inches(2.05), Inches(1.1), "Frontend", "Next.js 16", CYAN, NAVY)
abox = archbox(Inches(2.95), yc, Inches(2.05), Inches(1.1), "API Backend", "NestJS", TEAL)
obox = archbox(Inches(5.30), yc, Inches(2.45), Inches(1.1), "Orquestrador\nde Agentes", "máquina de estados", TEAL_DK)
connector(s, Inches(2.65), yc + Inches(0.55), Inches(2.95), yc + Inches(0.55))
connector(s, Inches(5.00), yc + Inches(0.55), Inches(5.30), yc + Inches(0.55))
# data sources column
srcs = [("RAG service", "/retrieve /generate /evaluate"), ("ChromaDB", "chunks vetoriais"),
        ("PostgreSQL", "dados + traces"), ("Ollama", "LLM + embeddings")]
sx = Inches(8.55); sw = Inches(4.15); shh = Inches(0.78); sy0 = Inches(2.05); sgap = Inches(0.20)
for i, (t, d) in enumerate(srcs):
    yy = sy0 + i*(shh+sgap)
    archbox(sx, yy, sw, shh, t, d, CARD, NAVY)
    rect(s, sx, yy, Inches(0.09), shh, [TEAL, CYAN, TEAL, CYAN][i], rounded=False)
    connector(s, Inches(7.75), yc + Inches(0.55), sx, yy + shh/2, color=SLATE_LT, width=Pt(1.5))
# response back
rbox = archbox(Inches(5.30), Inches(5.05), Inches(2.45), Inches(0.9), "Resposta ao usuário", "com fontes", GREEN)
connector(s, Inches(6.52), yc + Inches(1.1), Inches(6.52), Inches(5.05), color=GREEN, width=Pt(2))
# stack legend
txt(s, Inches(0.6), Inches(6.15), Inches(12), Inches(1.2),
    [P("Stack do MVP:  ", 12.5, TEAL, True) +
     [("Next.js 16  ·  NestJS/Fastify  ·  PostgreSQL/Drizzle  ·  ChromaDB  ·  "
       "Fastify RAG service  ·  Ollama local  ·  phi3:mini  ·  nomic-embed-text",
       12.5, SLATE, False, False)]], space_after=0)
footer(s, 7)
notes(s, "Esta é a visão geral. O Frontend em Next.js conversa com a API backend em "
         "NestJS, que funciona como fronteira de segurança: autentica, resolve a "
         "clínica ativa e fixa o escopo do paciente. Dentro da API, o Orquestrador de "
         "Agentes controla a execução. Ele chama o serviço RAG por endpoints "
         "primitivos, consulta Postgres para dados estruturados e traces, usa ChromaDB "
         "como índice vetorial e Ollama para embeddings e geração. A resposta volta "
         "por SSE com fontes, métricas e o identificador do agent_run.")

# ============================================================================
# SLIDE 8 — Fluxo de ingestão
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Pipeline", "Fluxo de ingestão dos documentos e anamneses",
       "Como prontuários e anotações viram conhecimento pesquisável.")
ing = [("Upload ou\nsnapshot", "📄", CYAN), ("API valida\nescopo", "🔐", TEAL),
       ("Parser", "🔧", TEAL), ("Chunking", "✂", TEAL),
       ("Embeddings", "🔢", TEAL_DK), ("ChromaDB", "🗃", GREEN)]
n = len(ing); bw = Inches(1.78); bh = Inches(1.3); y = Inches(2.85)
total = n*bw + (n-1)*Inches(0.27); x0 = (SW-total)/2
for i, (t, ic, col) in enumerate(ing):
    x = x0 + i*(bw+Inches(0.27))
    b = rect(s, x, y, bw, bh, CARD, rounded=True, shadow=True)
    rect(s, x, y, bw, Inches(0.12), col)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = ic; r.font.size = Pt(22); r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = t; r2.font.size = Pt(12.5); r2.font.bold = True
    r2.font.color.rgb = NAVY; r2.font.name = FONT
    if i < n-1:
        connector(s, x+bw, y+bh/2, x+bw+Inches(0.27), y+bh/2)
# metadata note box
mb = rect(s, Inches(3.4), Inches(4.7), Inches(6.5), Inches(1.0), CARD_ALT, rounded=True, shadow=True)
tf = mb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.3)
p = tf.paragraphs[0]; r = p.add_run(); r.text = "🗄  Metadados no PostgreSQL"
r.font.size = Pt(13.5); r.font.bold = True; r.font.color.rgb = TEAL_DK; r.font.name = FONT
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "patient_documents, anamneses, document_versions e status de ingestão preservam a fonte autoritativa."
r2.font.size = Pt(12); r2.font.color.rgb = SLATE; r2.font.name = FONT
connector(s, Inches(x0 + 5*(bw+Inches(0.27)) + bw/2), y+bh, Inches(6.6), Inches(4.7),
          color=SLATE_LT, width=Pt(1.5))
footer(s, 8)
notes(s, "A ingestão começa no backend, não no serviço RAG. A API valida o usuário, a "
         "clínica e o paciente, grava o arquivo ou snapshot de anamnese, cria o estado "
         "de ingestão e só então chama o RAG. O RAG parseia, faz chunking, gera "
         "embeddings com Ollama e grava no ChromaDB com metadados do paciente e da "
         "fonte. O Postgres continua sendo a fonte autoritativa: documentos, "
         "anamneses, versões e status ficam lá.")

# ============================================================================
# SLIDE 9 — Fluxo de pergunta no chat
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Exemplo prático", "Fluxo de uma pergunta no chat",
       None)
# question bubble
qb = rect(s, Inches(0.6), Inches(1.75), Inches(8.0), Inches(0.7), NAVY, rounded=True, shadow=True)
tf = qb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.3)
p = tf.paragraphs[0]; r = p.add_run()
r.text = "“Qual foi a última queixa de dor do paciente João?”"
r.font.size = Pt(15); r.font.italic = True; r.font.color.rgb = WHITE; r.font.name = FONT
flow = [
    ("Sessão validada", "JWT + X-Clinic-Id + patientId", "🔐"),
    ("IntentRouter", "classifica knowledge_base_search", "🧭"),
    ("Planner/Rewriter", "gera consultas normalizadas", "🧩"),
    ("RetrieverTool", "ChromaDB + Postgres fallback", "🔎"),
    ("Evaluate/Verify", "contexto suficiente + groundedness", "✔"),
]
y0 = Inches(2.75); rh = Inches(0.78); gap = Inches(0.12); x = Inches(0.6); w = Inches(8.0)
for i, (t, d, ic) in enumerate(flow):
    yy = y0 + i*(rh+gap)
    b = rect(s, x, yy, w, rh, CARD, rounded=True, shadow=True)
    rect(s, x, yy, Inches(0.09), rh, TEAL)
    badge = rect(s, x+Inches(0.22), yy+Inches(0.14), Inches(0.5), Inches(0.5), TEAL, rounded=True)
    bt = badge.text_frame; bt.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = bt.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run(); br.text = ic; br.font.size = Pt(17); br.font.color.rgb = WHITE; br.font.name = FONT
    txt(s, x+Inches(0.95), yy+Inches(0.08), Inches(2.6), rh,
        [P(t, 14, NAVY, True)], anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    txt(s, x+Inches(3.5), yy+Inches(0.08), w-Inches(3.7), rh,
        [P(d, 12.5, SLATE)], anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    num = txt(s, x-Inches(0.0), yy, Inches(0.0), rh, [P("", 1, SLATE)], space_after=0)
    if i < len(flow)-1:
        connector(s, x+Inches(0.47), yy+rh, x+Inches(0.47), yy+rh+gap, color=SLATE_LT, width=Pt(1.5))
# answer card on right
ab = rect(s, Inches(9.0), Inches(2.75), Inches(3.7), Inches(4.0), GREEN, rounded=True, shadow=True)
tf = ab.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.28); tf.margin_right = Inches(0.25)
tf.margin_top = Inches(0.25)
p = tf.paragraphs[0]; r = p.add_run(); r.text = "✅ Resposta fundamentada"
r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
p2 = tf.add_paragraph(); p2.space_before = Pt(8); r2 = p2.add_run()
r2.text = ("“A última queixa de dor de João foi em 12/03, dor no molar inferior "
           "direito, registrada na anamnese.”")
r2.font.size = Pt(13); r2.font.color.rgb = WHITE; r2.font.name = FONT
p3 = tf.add_paragraph(); p3.space_before = Pt(12); r3 = p3.add_run()
r3.text = "Fontes + métricas + agentRunId para inspecionar o trace"
r3.font.size = Pt(11); r3.font.italic = True; r3.font.color.rgb = CARD_ALT; r3.font.name = FONT
footer(s, 9)
notes(s, "No fluxo atual, o paciente não é escolhido pelo LLM: ele já está preso à "
         "sessão de chat validada pela API. A pergunta entra no endpoint agentic, o "
         "roteador classifica a intenção, o planner e o rewriter preparam consultas, "
         "o RetrieverTool busca no ChromaDB e pode complementar com Postgres, e o "
         "evaluator/verifier controlam contexto e fundamentação. A UI recebe tokens, "
         "fontes, métricas e o agentRunId, que permite abrir o trace completo.")

# ============================================================================
# SLIDE 10 — Agentes do sistema (tabela)
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Componentes", "Os agentes do sistema",
       "Cada agente tem uma responsabilidade única e bem definida.")
rows = [
    ["🧭  AgentOrchestrator", "Coordena a máquina de estados, orçamento e fallbacks"],
    ["🧠  IntentRouter", "Classifica direct_answer, RAG, multi-step, action_request ou unsupported"],
    ["🧩  QueryPlanner/Rewriter", "Decompõe a pergunta e gera consultas normalizadas"],
    ["🔎  RetrieverTool", "Consulta /v1/retrieve e pode usar Postgres como fallback/suplemento"],
    ["✔  ContextEvaluator", "Decide suficiência e sugere retry no loop CRAG"],
    ["✍  AnswerGenerator", "Gera resposta streamada com contexto recuperado"],
    ["🛡  AnswerVerifier", "Verifica groundedness e pode fazer downgrade seguro"],
    ["🗄  AgentTracing", "Persiste runs, steps, chunks, tools e avaliações"],
]
table(s, Inches(0.6), Inches(2.35), Inches(12.1),
      ["Agente", "Responsabilidade"], rows,
      col_widths=[Inches(3.5), Inches(8.6)], row_h=Inches(0.46), fs=11.5, head_fs=12.5)
footer(s, 10)
notes(s, "Esta tabela troca nomes genéricos por componentes reais do código. O "
         "AgentOrchestrator coordena a máquina de estados; o IntentRouter classifica "
         "a intenção; planner e rewriter estruturam as consultas; RetrieverTool chama "
         "o RAG e pode usar Postgres; ContextEvaluator decide se o contexto basta; "
         "AnswerGenerator gera a resposta; AnswerVerifier verifica groundedness; e "
         "AgentTracing persiste tudo para inspeção.")

# ============================================================================
# SLIDE 11 — Como o agente foi elaborado (anatomia do módulo)
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Engenharia", "Como o agente foi elaborado",
       "Módulo NestJS (src/agent) onde cada agente é um service injetável, "
       "coordenado por uma máquina de estados.")
rows = [
    ["orchestrator/", "AgentOrchestrator", "Máquina de estados / loop de controle"],
    ["router/", "IntentRouterService", "Classifica a intenção em 6 classes + confiança"],
    ["planner/", "QueryRewriter · QueryPlanner", "Normaliza, expande e decompõe a pergunta"],
    ["tools/", "ToolRegistry · RetrieverTool", "RAG, Postgres fallback e ações tipadas"],
    ["evaluators/", "ContextEvaluatorService", "Decide se o contexto recuperado é suficiente"],
    ["generators/", "AnswerGeneratorService", "Geração streamada (SSE) com citações inline"],
    ["verifiers/", "AnswerVerifierService", "Checa fundamentação (groundedness) das claims"],
    ["tracing/", "AgentTracingService", "Registra runs, steps, chunks, tools e evaluations"],
    ["memory/", "MemoryService", "Resumo da sessão para contexto futuro"],
]
table(s, Inches(0.6), Inches(2.2), Inches(12.1),
      ["Módulo (src/agent)", "Service (agente)", "Papel técnico"], rows,
      col_widths=[Inches(2.9), Inches(3.9), Inches(5.3)],
      row_h=Inches(0.46), fs=11.4, head_fs=12.5)
txt(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.6),
    [P("Princípio: ", 12.5, TEAL, True) +
     [("responsabilidade única por agente + injeção de dependência + feature flags "
       "para ligar/desligar cada capacidade de forma incremental.",
       12.5, SLATE, False, False)]], space_after=0)
footer(s, 11)
notes(s, "Tecnicamente, o agente é um módulo NestJS com serviços injetáveis. A parte "
         "importante para estudar é que o orquestrador não mistura responsabilidades: "
         "roteamento, planejamento, reescrita, recuperação, avaliação, geração, "
         "verificação, ferramentas, tracing e memória são componentes separados. Isso "
         "facilita testes, troca de modelos, feature flags e inspeção acadêmica do "
         "comportamento.")

# ============================================================================
# SLIDE 12 — O loop de orquestração + orçamento
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Controle", "O loop de orquestração",
       "Máquina de estados determinística com correção iterativa e orçamento (budget).")
loop_steps = [("Sanitizar", "🧼", CYAN), ("Rotear", "🧭", TEAL),
              ("Recuperar", "🔎", TEAL), ("Avaliar", "✔", TEAL),
              ("Gerar", "🧠", TEAL_DK), ("Verificar", "🛡", TEAL_DK),
              ("Persistir", "🗄", GREEN)]
n = len(loop_steps); bw = Inches(1.6); bh = Inches(1.15); y = Inches(3.25)
gapx = Inches(0.22)
total = n*bw + (n-1)*gapx; x0 = (SW-total)/2
centers = []
for i, (t, ic, col) in enumerate(loop_steps):
    x = x0 + i*(bw+gapx)
    centers.append(x + bw/2)
    b = rect(s, x, y, bw, bh, CARD, rounded=True, shadow=True)
    rect(s, x, y, bw, Inches(0.11), col)
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = ic; r.font.size = Pt(20); r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = t; r2.font.size = Pt(11.5); r2.font.bold = True
    r2.font.color.rgb = NAVY; r2.font.name = FONT
    if i < n-1:
        connector(s, x+bw, y+bh/2, x+bw+gapx, y+bh/2)
# feedback loop: Avaliar (idx 3) -> Recuperar (idx 2)
yline = y - Inches(0.55)
connector(s, centers[3], y, centers[3], yline, color=AMBER, width=Pt(1.6), arrow=False)
connector(s, centers[3], yline, centers[2], yline, color=AMBER, width=Pt(1.6), arrow=False)
connector(s, centers[2], yline, centers[2], y, color=AMBER, width=Pt(1.8), arrow=True)
txt(s, centers[2] - Inches(1.3), yline - Inches(0.42), Inches(3.6), Inches(0.4),
    [P("contexto insuficiente → reescreve a query", 11, AMBER, True)],
    align=PP_ALIGN.CENTER, space_after=0)
# budget card
bcard = rect(s, Inches(1.4), Inches(5.05), Inches(10.5), Inches(1.55), CARD_ALT,
             rounded=True, shadow=True)
tf = bcard.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.35); tf.margin_right = Inches(0.3)
p = tf.paragraphs[0]; r = p.add_run(); r.text = "⏱  Orçamento por execução (budget)"
r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = TEAL_DK; r.font.name = FONT
p2 = tf.add_paragraph(); p2.space_before = Pt(6); r2 = p2.add_run()
r2.text = ("maxLlmCalls (3)  ·  maxRetrievalAttempts  ·  maxWallClockMs (15 s) — "
           "limitam custo e latência; ao estourar, o loop é cortado e o sistema cai "
           "em um fallback seguro de “evidência insuficiente”.")
r2.font.size = Pt(12.5); r2.font.color.rgb = SLATE; r2.font.name = FONT
footer(s, 12)
notes(s, "O coração do agente é um loop determinístico, não uma cadeia mágica. A "
         "pergunta é primeiro sanitizada; o orquestrador roteia a intenção; recupera "
         "contexto; e o avaliador decide se é suficiente. Se não for, há a correção "
         "iterativa: ele reescreve a query e tenta de novo — essa é a seta laranja de "
         "retorno. Quando o contexto basta, gera a resposta, verifica a fundamentação "
         "e persiste. Tudo isso roda sob um orçamento explícito: número máximo de "
         "chamadas ao LLM, tentativas de recuperação e tempo de parede. Se o orçamento "
         "estoura, o loop é cortado e caímos num fallback seguro, evitando custo e "
         "latência descontrolados.")

# ============================================================================
# SLIDE 13 — Fundamentação, validação e segurança
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Confiabilidade", "Fundamentação, validação e segurança",
       "Os mecanismos técnicos que tornam a resposta confiável e o sistema seguro.")
mech = [
    ("ContextEvaluator", "Retorna { sufficient, missing[], suggestedQuery } e dispara o loop corretivo.", "✔", TEAL),
    ("AnswerVerifier", "Groundedness + citações; pode fazer downgrade para evidência insuficiente.", "🛡", TEAL),
    ("RAG-Triad", "Context relevance · groundedness · answer relevance persistidos no chat.", "📊", CYAN),
    ("Trust boundary", "JWT → ClinicScope → PatientScope: isolamento por clínica e paciente.", "🔐", TEAL_DK),
    ("Agent traces", "agent_runs, steps, chunks, tools e evaluations auditáveis no UI.", "🗄", CYAN),
    ("Preview → Commit", "Toda ferramenta mutável exige confirmação humana explícita.", "🤝", TEAL),
]
cw = Inches(3.85); ch = Inches(1.7); gx = Inches(0.22); gy = Inches(0.22)
x0 = Inches(0.6); y0 = Inches(2.3)
for i, (t, d, ic, col) in enumerate(mech):
    cx = i % 3; cy = i // 3
    card(s, x0 + cx*(cw+gx), y0 + cy*(ch+gy), cw, ch, t, d, accent=col, icon=ic)
footer(s, 13)
notes(s, "Confiabilidade e segurança são tratadas como mecanismos concretos. O "
         "ContextEvaluator produz uma estrutura tipada para decidir se o contexto "
         "basta. O AnswerVerifier controla groundedness e citações, e pode fazer "
         "downgrade para evidência insuficiente. A RAG-Triad fica persistida por "
         "mensagem. A fronteira de confiança é a API NestJS, com JWT, escopo de "
         "clínica e escopo de paciente. Além disso, os traces ficam auditáveis no "
         "Postgres e na UI. Por fim, qualquer ferramenta mutável passa por preview e "
         "confirmação humana.")

# ============================================================================
# SLIDE 14 — Como o agente faz a diferença (técnico)
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Impacto", "Como o agente faz a diferença",
       "Cada mecanismo resolve uma limitação concreta do RAG linear.")
rows = [
    ["Loop recuperar ↔ avaliar", "Recuperação única, sem 2ª chance",
     "Recall maior; contexto corrigido em runtime"],
    ["IntentRouter + Planner", "Não decompõe perguntas",
     "Perguntas multi-etapa são atendidas"],
    ["RetrieverTool híbrido", "Só busca no vetor store",
     "ChromaDB + Postgres fallback/suplemento"],
    ["ToolRegistry tipado", "Não usa ferramentas",
     "Agendamentos via preview → commit"],
    ["Verifier + citações", "Não valida a resposta",
     "Menos alucinação; resposta rastreável"],
    ["Registro RAG no UI", "Caixa-preta, difícil estudar",
     "Passo a passo auditável por execução"],
]
cell_colors = {}
for i in range(len(rows)):
    cell_colors[(i, 1)] = RGBColor(0xFE, 0xF3, 0xEC)  # amber-ish (limitação)
    cell_colors[(i, 2)] = RGBColor(0xE9, 0xF9, 0xEF)  # green-ish (efeito)
table(s, Inches(0.6), Inches(2.3), Inches(12.1),
      ["Mecanismo do agente", "Limitação do RAG linear", "Efeito técnico"], rows,
      col_widths=[Inches(3.4), Inches(4.0), Inches(4.7)],
      row_h=Inches(0.72), cell_colors=cell_colors, fs=12, head_fs=12.5)
txt(s, Inches(0.6), Inches(6.55), Inches(12), Inches(0.6),
    [P("Em uma frase: ", 12.5, TEAL, True) +
     [("o agente troca um pipeline fixo por um controlador que planeja, recupera, "
       "valida e prova — com custo limitado e rastreável.",
       12.5, SLATE, False, False)]], space_after=0)
footer(s, 14)
notes(s, "Aqui o ponto é mostrar diferença prática. O loop recuperar-avaliar dá uma "
         "segunda chance ao contexto. O roteador e o planner permitem perguntas "
         "multi-etapa. O RetrieverTool híbrido evita depender exclusivamente do vetor "
         "store. O ToolRegistry habilita agendamentos com confirmação. O verifier e "
         "as citações reduzem alucinação. E o registro RAG no UI transforma uma "
         "resposta em objeto de estudo: dá para ver cada step, input, output, chunks, "
         "scores, latência e fallback.")

# ============================================================================
# SLIDE 15 — Escolhas técnicas: prós e contras
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Decisões", "Escolhas técnicas: prós e contras",
       "Comparações que guiaram a stack do MVP acadêmico.")
rows = [
    ["Banco vetorial",
     "ChromaDB: simples, metadata filter, separado do domínio",
     "pgvector: menos serviços, acopla vetores ao Postgres"],
    ["LLM",
     "Ollama local: privacidade e custo zero",
     "OpenAI/Claude: mais qualidade, dependência externa"],
    ["Backend",
     "NestVJS: produtivo, ecossistema TS",
     "Go: performance, menos boilerplate ORM"],
    ["Abordagem",
     "RAG legado: rápido e simples",
     "Agentic RAG: auditável, robusto, mais complexo"],
]
# fix typo
rows[2][1] = "NestJS: produtivo, ecossistema TS/Node"
cell_colors = {}
for i in range(len(rows)):
    cell_colors[(i, 1)] = RGBColor(0xE9, 0xF9, 0xEF)  # green-ish
    cell_colors[(i, 2)] = RGBColor(0xFE, 0xF3, 0xEC)  # amber-ish
table(s, Inches(0.6), Inches(2.35), Inches(12.1),
      ["Decisão", "Opção A — vantagem", "Opção B — trade-off"], rows,
      col_widths=[Inches(2.5), Inches(4.8), Inches(4.8)],
      row_h=Inches(0.82), cell_colors=cell_colors, fs=12, head_fs=12.5)
txt(s, Inches(0.6), Inches(6.45), Inches(12), Inches(0.6),
    [P("No estado atual: ", 12.5, TEAL, True) +
     [("ChromaDB é índice derivado; Postgres é fonte autoritativa; Ollama mantém "
       "execução local; Agentic RAG roda atrás de flags e budgets.",
       12.5, SLATE, False, False)]], space_after=0)
footer(s, 15)
notes(s, "Aqui eu atualizo as decisões para o estado real do projeto. O banco vetorial "
         "é ChromaDB, escolhido por simplicidade e filtro por metadados. Postgres "
         "continua sendo a fonte autoritativa. Ollama mantém modelos locais, o que "
         "favorece privacidade e custo zero em ambiente acadêmico. NestJS concentra a "
         "fronteira de segurança e a orquestração. O RAG legado continua existindo, "
         "mas o Agentic RAG entrega mais auditoria e robustez ao custo de mais "
         "complexidade.")

# ============================================================================
# SLIDE 16 — Vantagens do Agentic RAG
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Benefícios", "Vantagens do Agentic RAG",
       "Por que vale a pena a complexidade adicional.")
adv = [
    ("Melhor recuperação", "Busca iterativa e corretiva do contexto.", "🎯"),
    ("Múltiplas etapas", "Decompõe perguntas complexas em passos.", "🧩"),
    ("Menos alucinação", "Validação de evidência antes de responder.", "🛡"),
    ("Ferramentas dinâmicas", "Executa ações de agenda com confirmação humana.", "🛠"),
    ("Explicabilidade", "Mostra steps, métricas, chunks e fontes.", "🔍"),
    ("Tarefas complexas", "Adequado a fluxos de trabalho reais da clínica.", "⚙"),
]
cw = Inches(3.85); ch = Inches(1.5); gx = Inches(0.22); gy = Inches(0.20)
x0 = Inches(0.6); y0 = Inches(2.35)
for i, (t, d, ic) in enumerate(adv):
    cx = i % 3; cy = i // 3
    card(s, x0 + cx*(cw+gx), y0 + cy*(ch+gy), cw, ch, t, d, accent=GREEN, icon=ic)
footer(s, 16)
notes(s, "As vantagens do Agentic RAG vêm justamente de superar as limitações do RAG "
         "tradicional. A recuperação melhora porque é iterativa e corretiva. O sistema "
         "lida com perguntas de múltiplas etapas. A validação de evidência reduz "
         "alucinações. A escolha dinâmica de ferramentas permite executar ações de "
         "agendamento com confirmação humana. A explicabilidade aumenta porque o "
         "sistema mostra steps, métricas, chunks e fontes. No conjunto, fica muito mais "
         "adequado a tarefas reais e complexas de uma clínica.")

# ============================================================================
# SLIDE 17 — Desvantagens e riscos
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Riscos", "Desvantagens e riscos",
       "Trade-offs que precisam ser gerenciados conscientemente.")
risks = [
    ("Maior complexidade", "Mais componentes para construir e manter.", "🧱"),
    ("Mais custo de tokens", "Múltiplas chamadas ao LLM por pergunta.", "💸"),
    ("Maior latência", "Etapas extras aumentam o tempo de resposta.", "⏱"),
    ("PII em traces", "Payloads brutos exigem política de redação em produção.", "📉"),
    ("Controle de permissões", "Isolar dados por paciente é obrigatório.", "🔐"),
    ("Ações indevidas", "Risco de agentes agirem sem confirmação.", "🚫"),
]
cw = Inches(3.85); ch = Inches(1.5); gx = Inches(0.22); gy = Inches(0.20)
x0 = Inches(0.6); y0 = Inches(2.35)
for i, (t, d, ic) in enumerate(risks):
    cx = i % 3; cy = i // 3
    card(s, x0 + cx*(cw+gx), y0 + cy*(ch+gy), cw, ch, t, d, accent=RED, icon=ic)
footer(s, 17)
notes(s, "É preciso ser honesto sobre os riscos. O Agentic RAG é mais complexo de "
         "construir e manter. Cada pergunta pode gerar várias chamadas ao LLM, "
         "elevando custo de tokens e latência. Como os traces armazenam dados úteis "
         "para estudo, em produção seria necessário redigir PII com mais rigor. Como "
         "lidamos com dados clínicos, o controle de permissões e o isolamento por "
         "paciente são obrigatórios. E há o risco de um agente executar ações "
         "indevidas — por isso adotamos confirmação humana (preview → commit) para "
         "qualquer ação que altere dados.")

# ============================================================================
# SLIDE 18 — Escopo do MVP
# ============================================================================
s = slide(); bg(s, BG_LIGHT)
header(s, "Delimitação", "Escopo do MVP acadêmico",
       "O que entra e o que fica de fora, conscientemente.")
# Included panel
inc = rect(s, Inches(0.6), Inches(2.3), Inches(5.95), Inches(4.4), CARD, rounded=True, shadow=True)
rect(s, Inches(0.6), Inches(2.3), Inches(5.95), Inches(0.62), GREEN, rounded=False)
txt(s, Inches(0.9), Inches(2.42), Inches(5.4), Inches(0.4),
    [P("✅  Incluído no MVP", 16, WHITE, True)], space_after=0)
bullets(s, Inches(0.95), Inches(3.15), Inches(5.3), [
    "Cadastro de pacientes",
    "Anamnese e documentos",
    "Chat com Agentic RAG",
    "Busca ChromaDB + Postgres",
    "Ferramentas de agenda",
    "Registro RAG inspecionável",
], size=14.5, gap=12, marker_color=GREEN)
# Out panel
out = rect(s, Inches(6.78), Inches(2.3), Inches(5.95), Inches(4.4), CARD, rounded=True, shadow=True)
rect(s, Inches(6.78), Inches(2.3), Inches(5.95), Inches(0.62), RED, rounded=False)
txt(s, Inches(7.08), Inches(2.42), Inches(5.4), Inches(0.4),
    [P("🚫  Fora do MVP", 16, WHITE, True)], space_after=0)
bullets(s, Inches(7.13), Inches(3.15), Inches(5.3), [
    "Integração real com sistemas hospitalares",
    "Diagnóstico médico automatizado",
    "Ações automáticas sem confirmação",
    "Treinamento próprio de LLM",
    "Política completa de LGPD/produção",
], size=14.5, gap=14, marker_color=RED)
footer(s, 18)
notes(s, "Delimitar o escopo é essencial num trabalho acadêmico. Entram no MVP: "
         "cadastro de pacientes, anamnese, documentos, chat com Agentic RAG, busca "
         "semântica em ChromaDB combinada com Postgres, ferramentas de agenda e o "
         "registro RAG inspecionável. Ficam deliberadamente de fora: integração real "
         "com sistemas hospitalares, diagnóstico médico automatizado, ações sem "
         "confirmação humana, treinamento próprio de LLM e uma política completa de "
         "produção/LGPD. Esses limites mantêm o projeto viável e eticamente "
         "responsável.")

# ============================================================================
# SLIDE 19 — Conclusão
# ============================================================================
s = slide(); bg(s, NAVY)
d = rect(s, Inches(-1.5), Inches(4.5), Inches(7), Inches(7), TEAL_DK, rounded=True); d.rotation = 20
d2 = rect(s, Inches(9.5), Inches(-2), Inches(6), Inches(6), TEAL, rounded=True); d2.rotation = 20
rect(s, Inches(0.9), Inches(1.2), Inches(0.6), Inches(0.13), CYAN)
txt(s, Inches(0.9), Inches(1.42), Inches(8), Inches(0.4),
    [P("CONCLUSÃO", 14, CYAN, True)], space_after=0)
txt(s, Inches(0.86), Inches(2.0), Inches(11.5), Inches(2.6),
    [P("O Agentic RAG torna o CRM Dental", 32, WHITE, True),
     P("mais inteligente: o sistema não apenas ", 32, WHITE, True),
     [("busca", 32, CYAN, True, False)] +
     [(" — ele ", 32, WHITE, True, False)] +
     [("planeja, valida, usa ferramentas", 32, CYAN, True, False)],
     P("e deixa sua execução auditável.", 32, WHITE, True)],
    space_after=4, line_spacing=1.05)
# takeaways row
tks = ["Planeja a tarefa", "Valida evidência", "Usa ferramentas", "Registra o trace"]
bw = Inches(2.85); gap = Inches(0.25); x0 = Inches(0.9); y = Inches(5.55)
for i, t in enumerate(tks):
    pill(s, x0 + i*(bw+gap), y, bw, Inches(0.85), t,
         TEAL if i % 2 == 0 else CYAN, WHITE if i % 2 == 0 else NAVY, size=14)
notes(s, "Para concluir: o Agentic RAG torna o CRM Dental mais inteligente porque "
         "muda a natureza do sistema. Ele deixa de apenas buscar informação e passa a "
         "planejar a tarefa, validar a evidência, usar ferramentas quando necessário e "
         "registrar sua própria execução. Para um contexto clínico e acadêmico, onde "
         "confiança, rastreabilidade e estudo técnico são críticos, essa evolução faz "
         "diferença real — mesmo num MVP.")

# ============================================================================
# SLIDE 20 — Perguntas?
# ============================================================================
s = slide(); bg(s, NAVY)
d2 = rect(s, Inches(9.0), Inches(3.5), Inches(7), Inches(7), TEAL_DK, rounded=True); d2.rotation = 25
d3 = rect(s, Inches(10.5), Inches(0.0), Inches(2.4), Inches(2.4), CYAN, rounded=True); d3.rotation = 25
txt(s, Inches(0), Inches(2.7), SW, Inches(1.6),
    [P("Perguntas?", 60, WHITE, True)], align=PP_ALIGN.CENTER, space_after=0)
txt(s, Inches(0), Inches(4.4), SW, Inches(0.6),
    [P("Obrigado pela atenção.", 20, CYAN, False)], align=PP_ALIGN.CENTER, space_after=0)
txt(s, Inches(0), Inches(5.1), SW, Inches(0.5),
    [P("CRM Dental com Chat Inteligente usando Agentic RAG", 14, SLATE_LT, False)],
    align=PP_ALIGN.CENTER, space_after=0)
notes(s, "Encerro por aqui e abro para perguntas. Obrigado pela atenção. Estou à "
         "disposição para detalhar a arquitetura, as escolhas técnicas ou o escopo do "
         "MVP.")

# ----------------------------------------------------------------------------
out_path = "CRM-Dental-Agentic-RAG.pptx"
prs.save(out_path)
print("Apresentação salva em:", out_path, "| slides:", len(prs.slides._sldIdLst))
